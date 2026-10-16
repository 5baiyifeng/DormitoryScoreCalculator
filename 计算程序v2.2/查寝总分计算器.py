# -*- coding: utf-8 -*-
"""
宿委会查寝总分计算器 v2.2
------------------------------------------------------------
在 v2.0 基础上新增：等级（优秀/良好/合格/不合格）由程序生成，
不再依赖 Excel 公式。

功能：
  1. 选择"总表原件"Excel，程序自动在同目录生成带时间戳的副本，
     之后所有读写都只针对副本，绝不修改原件。
  2. 输入三位寝室号 n，在 F 列(tag) 中匹配行号 i，把本次得分写入 D 列第 i 行。
  3. 每个寝室录入最多 11 个"扣分绝对值"，总分 = 100 - 各项之和；
     输入 0 提前结束本次计算，填满 11 项自动结束。
  4. 写分数时同步写入 E 列"等级"（程序判级：<60 不合格，60~79 合格，
     80~89 良好，≥90 优秀）；副本里原有的等级公式会被删掉，避免出现
     公式结果覆盖程序结果的情况。
  5. 若 D 列第 i 行已有分数，弹窗确认是否覆盖（快捷键 Y / N），
     覆盖时分数和等级一起更新。
  6. 结算后等待用户按回车再进入下一个寝室（防连按、便于核对）；
     输入 -1 结束程序并保存副本。
"""

import os
import re
import sys
import shutil
import datetime
import tkinter as tk
from tkinter import filedialog, messagebox

import openpyxl

# ---------------------------------------------------------------- 常量配置
VERSION = "v2.2"
MAX_ITEMS = 11          # 纸质表共 11 个扣分词段
FULL_SCORE = 100        # 基准分
WAIT_ENTER = True       # 结算后等待手动按回车（True）；改为 False 则用 PAUSE_MS 自动继续
PAUSE_MS = 2000         # 自动继续模式下每次计算后的防连按暂停（毫秒）
D_FIRST_ROW = 2         # 数据起始行
TAG_COL = 6             # F 列为 tag 字段
SCORE_COL = 4           # D 列为分数列
GRADE_COL = 5           # E 列为等级列
END_CODE = "-1"         # 结束程序
HINT_TEXT = "记得看看有没有写上厕所 -3"

# ---------------- 免责声明（每次启动程序时都要先看并确认） ----------------
DISCLAIMER_TITLE = "免责声明"
AUTHOR_QQ = "3521228485"
DISCLAIMER_TEXT = (
    "【免责声明】\n\n"
    "本程序是作者第一次动手写的小工具，可能还存在没有发现的 Bug，\n"
    "计算结果仅供参考。\n\n"
    "请务必自行核实程序生成的表格内容（分数、等级、寝室号是否对得上），\n"
    "确认无误后再使用或上缴。\n\n"
    "如果发现程序出错、数字不对，或者有改进建议，\n"
    "欢迎加 QQ 反馈：%s\n\n"
    "点击“我已知晓”或直接按回车，开始使用程序。" % AUTHOR_QQ
)

# 等级分档：分数下限（含）→ 等级名称，从高到低排列
GRADE_LEVELS = [
    (90, "优秀"),
    (80, "良好"),
    (60, "合格"),
    (0, "不合格"),
]

FONT_UI = ("Microsoft YaHei UI", 11)
FONT_BIG = ("Microsoft YaHei UI", 16, "bold")
FONT_MONO = ("Consolas", 14)


# ---------------------------------------------------------------- 工具函数
def norm_tag(value):
    """把 tag 值统一成可比较的字符串：数字 201 与文本 '201' 都能对上。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def cell_has_value(value):
    """判断单元格是否已有内容。"""
    if value is None:
        return False
    if isinstance(value, str):
        return value.strip() != ""
    return True


def grade_of(score):
    """按分数算出等级；分数为空则返回 None（不写等级）。"""
    if score is None:
        return None
    try:
        value = float(score)
    except (TypeError, ValueError):
        return None
    for lower, name in GRADE_LEVELS:
        if value >= lower:
            return name
    return None


def make_copy_path(src_path):
    """在原件同一文件夹内生成带时间戳的副本路径（不覆盖已有文件）。"""
    folder = os.path.dirname(os.path.abspath(src_path))
    stem, ext = os.path.splitext(os.path.basename(src_path))
    if not ext:
        ext = ".xlsx"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    candidate = os.path.join(folder, "%s - 副本_%s%s" % (stem, stamp, ext))
    seq = 2
    while os.path.exists(candidate):
        candidate = os.path.join(folder, "%s - 副本_%s(%d)%s" % (stem, stamp, seq, ext))
        seq += 1
    return candidate


def load_tag_map(src_path):
    """
    读取原件（只读，用于取缓存值），返回 (工作表名, {归一化tag: 行号})。
    tag 为公式时取 Excel 上次保存的缓存值。
    """
    wb = openpyxl.load_workbook(src_path, data_only=True, read_only=True)
    result = None
    for ws in wb.worksheets:
        if ws.max_column is None or ws.max_column < TAG_COL:
            continue
        tag_map = {}
        row = D_FIRST_ROW
        empty_streak = 0
        while row <= (ws.max_row or 0):
            key = norm_tag(ws.cell(row, TAG_COL).value)
            if key:
                empty_streak = 0
                if key not in tag_map:
                    tag_map[key] = row
            else:
                empty_streak += 1
                if empty_streak >= 20:      # 连续空白视为表末
                    break
            row += 1
        if tag_map:
            result = (ws.title, tag_map)
            break
    wb.close()
    if result is None:
        raise ValueError(
            "没有在所选表格的 F 列（tag）中找到任何寝室号。\n"
            "请确认所选文件是含 tag 列的总表原件。"
        )
    return result


# ---------------------------------------------------------------- 主窗口
class ScoreApp:
    def __init__(self, root, src_path):
        self.root = root
        self.src_path = os.path.abspath(src_path)

        self.entries = []
        self.items = []          # 已录入的扣分
        self.locked = False      # 结算/等待确认期间锁定扣分格
        self.awaiting_enter = False   # 已结算，等待用户按回车进入下一个寝室
        self.pending_after = None
        self.current_tag = None
        self.row_by_tag = {}
        self.sheet_title = ""
        self.wb = None
        self.ws = None
        self.formula_rows_cleaned = 0   # 被程序改写过的等级公式行数

        self._build_ui()
        if not self._prepare_files():
            self.root.destroy()
            return
        self.room_entry.focus_set()

    # ------------------------------------------------------------ 界面
    def _build_ui(self):
        self.root.title("宿委会查寝总分计算器 %s" % VERSION)
        self.root.resizable(False, False)
        self.root.protocol("WM_DELETE_WINDOW", self.on_quit)

        top = tk.Frame(self.root, padx=12, pady=10)
        top.pack(fill="x")

        tk.Label(top, text="寝室号（三位数，输入 -1 结束程序）：",
                 font=FONT_UI).grid(row=0, column=0, sticky="w")
        self.room_var = tk.StringVar()
        self.room_entry = tk.Entry(top, textvariable=self.room_var, font=FONT_MONO,
                                   width=8, justify="center")
        self.room_entry.grid(row=0, column=1, sticky="w", padx=6)
        self.room_entry.bind("<Return>", self.on_room_enter)

        self.room_hint = tk.Label(top, text="", font=FONT_UI, fg="#0a6")
        self.room_hint.grid(row=0, column=2, sticky="w", padx=6)

        tk.Label(self.root, text=HINT_TEXT, font=FONT_UI,
                 fg="#c00").pack(anchor="w", padx=12)

        grid = tk.LabelFrame(self.root, text=" 扣分录入（输入 0 提前结束，填满 11 项自动结束） ",
                             font=FONT_UI, padx=10, pady=8)
        grid.pack(fill="x", padx=12, pady=8)

        for i in range(MAX_ITEMS):
            r, c = divmod(i, 4)
            tk.Label(grid, text="%d." % (i + 1), font=FONT_UI).grid(
                row=r, column=c * 2, sticky="e", padx=(6, 2), pady=4)
            var = tk.StringVar()
            vcmd = (self.root.register(self._validate_item), "%P", "%s", "%S",
                    str(i), id(var))
            ent = tk.Entry(grid, textvariable=var, font=FONT_MONO, width=5,
                           justify="center", state="disabled",
                           validate="key", validatecommand=vcmd)
            ent.grid(row=r, column=c * 2 + 1, sticky="w", padx=(0, 10), pady=4)
            ent.bind("<Return>", lambda e, idx=i: self.on_item_enter(idx))
            self.entries.append((var, ent))

        bottom = tk.Frame(self.root, padx=12, pady=6)
        bottom.pack(fill="x")
        self.score_var = tk.StringVar(value="本次得分：--")
        tk.Label(bottom, textvariable=self.score_var, font=FONT_BIG,
                 fg="#06c").pack(side="left")
        self.grade_var = tk.StringVar(value="")
        tk.Label(bottom, textvariable=self.grade_var, font=FONT_BIG,
                 fg="#0a6").pack(side="left", padx=12)
        tk.Button(bottom, text="结束并保存 (-1)", font=FONT_UI,
                  command=self.on_quit).pack(side="right")

        self.status_var = tk.StringVar(value="正在准备文件……")
        tk.Label(self.root, textvariable=self.status_var, font=("Microsoft YaHei UI", 9),
                 fg="#555", anchor="w", justify="left", wraplength=560).pack(
            fill="x", padx=12, pady=(0, 10))

    # ------------------------------------------------------------ 文件准备
    def _prepare_files(self):
        try:
            self.sheet_title, self.row_by_tag = load_tag_map(self.src_path)
        except Exception as exc:
            messagebox.showerror("读取失败", "读取原件失败：\n%s" % exc)
            return False

        self.copy_path = make_copy_path(self.src_path)
        try:
            shutil.copy2(self.src_path, self.copy_path)
        except Exception as exc:
            messagebox.showerror("复制失败", "无法创建副本：\n%s" % exc)
            return False

        try:
            self.wb = openpyxl.load_workbook(self.copy_path)
            self.ws = self.wb[self.sheet_title]
        except Exception as exc:
            messagebox.showerror("打开副本失败", str(exc))
            return False

        # 等级改由程序生成：把副本里原有的等级公式就地替换成程序算出的等级
        self._take_over_grades()

        self._show_status()
        return True

    def _take_over_grades(self):
        """把副本 E 列的等级公式改成程序算出的等级文字（原公式不再保留）。"""
        ws = self.ws
        changed = 0
        last_row = max(self.row_by_tag.values()) if self.row_by_tag else D_FIRST_ROW
        for row in range(D_FIRST_ROW, last_row + 1):
            cell = ws.cell(row, GRADE_COL)
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                score = ws.cell(row, SCORE_COL).value
                cell.value = grade_of(score)      # 没有分数就留空
                changed += 1
        self.formula_rows_cleaned = changed
        if changed:
            self._save_copy()

    def _save_copy(self):
        try:
            self.wb.save(self.copy_path)
            return True
        except Exception as exc:
            messagebox.showerror("保存失败", "写入副本失败：\n%s" % exc)
            return False

    def _show_status(self):
        parts = ["宿舍委会查寝总分计算器 %s" % VERSION,
                 "原件：%s" % self.src_path,
                 "副本：%s" % self.copy_path,
                 "已载入 %d 个寝室号。" % len(self.row_by_tag)]
        if self.formula_rows_cleaned:
            parts.append("等级列原有的 %d 行公式已由程序接管（改为按分数生成等级文字）。"
                         % self.formula_rows_cleaned)
        else:
            parts.append("等级由程序按分数生成（<60 不合格，60~79 合格，80~89 良好，≥90 优秀）。")
        self.status_var.set("\n".join(parts))

    # ------------------------------------------------------------ 交互逻辑
    def _focus_score_entry(self):
        """把焦点交给第一个扣分输入框。"""
        self._unlock_entries()
        self.entries[0][1].focus_set()

    def _lock_entries(self):
        self.locked = True
        for _, ent in self.entries:
            ent.configure(state="disabled")

    def _unlock_entries(self):
        self.locked = False
        for _, ent in self.entries:
            ent.configure(state="normal")

    def on_room_enter(self, _event=None):
        """寝室号框按回车：结算后的确认、或者开始录入一个新寝室。"""
        text = self.room_var.get().strip()

        # 结算完成后等待确认：-1 退出，空回车继续，直接输入新寝室号也可以
        if self.awaiting_enter:
            if text == END_CODE:
                self.on_quit()
                return
            if text == "":
                self.next_room()
                return
        elif self.locked:
            return

        if text == END_CODE:
            self.on_quit()
            return
        if not re.fullmatch(r"\d{3}", text):
            messagebox.showwarning("输入有误", "寝室号必须是三位整数，例如 405。")
            self.room_entry.select_range(0, "end")
            return
        if text not in self.row_by_tag:
            messagebox.showwarning(
                "未找到寝室",
                "在 F 列（tag）中没有找到 %s。\n请核对寝室号，或检查所选原件是否正确。" % text)
            self.room_entry.select_range(0, "end")
            return

        self.awaiting_enter = False
        self.current_tag = text
        self.room_hint.configure(text="当前寝室：%s" % text)
        self.score_var.set("本次得分：--")
        self.grade_var.set("")
        for var, _ in self.entries:
            var.set("")
        self.items = []
        self._focus_score_entry()

    def _validate_item(self, proposed, previous, inserted, idx, var_key):
        """扣分框只允许敲入 0~999 的整数；结算统一由回车触发，避免误判。"""
        if proposed != "" and not re.fullmatch(r"\d{0,3}", proposed):
            return False
        if self.locked:
            # 结算/等待确认期间：只允许清空（换寝室时清空），不接受新输入
            return proposed == ""
        return True

    def on_item_enter(self, idx):
        """在第 idx 个扣分框按回车：记录这一格的数字，然后跳到下一格。

        填 0 表示本次录入到此为止；第 11 格回车则自动结算。
        """
        if self.locked:
            return
        if idx != len(self.items):
            # 前面的格子还没填完就被跳过了
            messagebox.showwarning(
                "顺序不对",
                "请从第 %d 格开始按顺序填。\n没有扣分的格子可以填 0 表示到此为止。"
                % (len(self.items) + 1))
            if len(self.items) < MAX_ITEMS:
                self.entries[len(self.items)][1].focus_set()
            return
        raw = self.entries[idx][0].get().strip()
        if raw == "":
            messagebox.showwarning("还没填完", "第 %d 格还没有填数字。\n"
                                                "这一项没扣分就填 0（表示到此为止）。" % (idx + 1))
            return
        self._commit_value(int(raw))

    def _commit_value(self, value):
        """记录一项扣分并累计总分；值为 0 表示本次到此为止，满 11 项自动结算。"""
        self.items.append(value)
        next_idx = len(self.items)
        if value == 0:
            self.finish_round("结束标志 0")
        elif next_idx >= MAX_ITEMS:
            self.finish_round("已填满 %d 项" % MAX_ITEMS)
        elif next_idx < MAX_ITEMS:
            self.entries[next_idx][1].focus_set()

    # ------------------------------------------------------------ 结算
    def finish_round(self, reason):
        """结算：写分数与等级，进入“等待确认”状态，由用户按回车进入下一个寝室。"""
        self._lock_entries()
        self.awaiting_enter = True
        score = FULL_SCORE - sum(self.items)
        filled = len([v for v in self.items if v != 0])
        grade = grade_of(score)
        self.score_var.set("本次得分：%d" % score)
        self.grade_var.set("等级：%s" % grade)
        self.status_var.set("录入 %d 个扣分（%s），得分 %d，等级 %s。"
                            % (filled, reason, score, grade))

        self.apply_score(score, grade)

        if WAIT_ENTER:
            # 手动确认：把光标放到寝室号框，直接按回车即可继续；也可以直接输入下一个寝室号
            self.status_var.set(self.status_var.get() +
                                "\n请核对上面显示的分数与等级，按回车进入下一个寝室。")
            self.room_var.set("")
            self.room_hint.configure(text="按回车继续")
            self.room_entry.focus_set()
        else:
            self.pending_after = self.root.after(PAUSE_MS, self.next_room)

    def apply_score(self, score, grade):
        """把分数写进 D 列、等级写进 E 列；已有分数时先确认是否覆盖。"""
        tag = self.current_tag
        row = self.row_by_tag.get(tag)
        if row is None:
            messagebox.showerror("内部错误", "找不到寝室 %s 对应的行，本次结果未写入。" % tag)
            return
        old = self.ws.cell(row, SCORE_COL).value
        if cell_has_value(old):
            old_grade = self.ws.cell(row, GRADE_COL).value
            ok = messagebox.askyesno(
                "覆盖确认",
                "该寝室已有分数，是否覆盖？\n\n是（Y）或否（N）\n\n"
                "当前：%s 分（等级 %s）\n本次：%d 分（等级 %s）"
                % (old, old_grade if cell_has_value(old_grade) else "—", score, grade),
                default=messagebox.NO)
            if not ok:
                self.status_var.set("寝室 %s 保留原有分数 %s，本次结果已丢弃。" % (tag, old))
                return
        self.ws.cell(row, SCORE_COL).value = score
        self.ws.cell(row, GRADE_COL).value = grade      # 等级由程序生成，不再用公式
        if self._save_copy():
            self.status_var.set("已写入：寝室 %s（第 %d 行）分数 %d、等级 %s，副本已保存。"
                                % (tag, row, score, grade))

    def next_room(self):
        """回到寝室号输入，准备量下一个寝室。"""
        self.pending_after = None
        self.awaiting_enter = False
        self.room_var.set("")
        self.room_hint.configure(text="")
        self.score_var.set("本次得分：--")
        self.grade_var.set("")
        for var, _ in self.entries:
            var.set("")
        self.items = []
        self.current_tag = None
        self._unlock_entries()
        for _, ent in self.entries:
            ent.configure(state="disabled")
        self._show_status()
        self.room_entry.focus_set()

    # ------------------------------------------------------------ 退出
    def on_quit(self):
        if self.pending_after is not None:
            self.root.after_cancel(self.pending_after)
            self.pending_after = None
        try:
            if self.wb is not None:
                self.wb.save(self.copy_path)
        except Exception as exc:
            messagebox.showerror("保存失败", "退出前保存副本失败：\n%s" % exc)
            return
        messagebox.showinfo("已结束",
                            "程序结束。\n\n最终副本已生成：\n%s" % getattr(self, "copy_path", ""))
        self.root.destroy()


# ---------------------------------------------------------------- 入口
def show_disclaimer(root):
    """每次启动都要先显示免责声明：放在屏幕正中，必须按回车或点按钮才能继续。"""
    dialog = tk.Toplevel(root)
    dialog.title(DISCLAIMER_TITLE)
    dialog.resizable(False, False)

    frame = tk.Frame(dialog, padx=18, pady=14)
    frame.pack(fill="both", expand=True)

    tk.Label(frame, text=DISCLAIMER_TEXT, font=("Microsoft YaHei UI", 11),
             fg="#cc0000", justify="left", anchor="w").pack(fill="both", expand=True)

    button = tk.Button(frame, text="我已知晓（回车）", font=("Microsoft YaHei UI", 11),
                       width=18, height=1, command=dialog.destroy)
    button.pack(pady=(14, 0))

    def close(_event=None):
        try:
            dialog.destroy()
        except Exception:
            pass

    dialog.bind("<Return>", close)
    dialog.bind("<space>", close)
    dialog.bind("<Escape>", close)
    dialog.protocol("WM_DELETE_WINDOW", close)   # 点 × 也算确认过一次

    # 放在屏幕正中（先算尺寸再定位，确保看得见）
    dialog.update_idletasks()
    w, h = dialog.winfo_reqwidth(), dialog.winfo_reqheight()
    x = max((dialog.winfo_screenwidth() - w) // 2, 0)
    y = max((dialog.winfo_screenheight() - h) // 2, 0)
    dialog.geometry("%dx%d+%d+%d" % (w, h, x, y))
    dialog.lift()
    dialog.focus_force()
    try:
        dialog.attributes("-topmost", True)
    except Exception:
        pass
    button.focus_set()

    dialog.grab_set()
    dialog.wait_window()          # 阻塞到用户确认
    try:
        dialog.attributes("-topmost", False)
    except Exception:
        pass


def choose_file():
    """弹出文件选择框，默认从程序所在目录开始找。"""
    start = os.path.dirname(os.path.abspath(sys.executable if getattr(sys, "frozen", False)
                                            else __file__))
    return filedialog.askopenfilename(
        title="请选择查寝总表原件（Excel）",
        initialdir=start,
        filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")])


def write_crash_log(exc_text):
    """把启动异常写到 exe 同目录的 出错日志.txt，避免窗口版程序“无声消失”。"""
    try:
        if getattr(sys, "frozen", False):
            folder = os.path.dirname(os.path.abspath(sys.executable))
        else:
            folder = os.path.dirname(os.path.abspath(__file__))
        path = os.path.join(folder, "出错日志.txt")
        with open(path, "a", encoding="utf-8") as f:
            f.write("=" * 60 + "\n")
            f.write("时间：%s\n" % datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
            f.write("程序：查寝总分计算器 %s\n\n" % VERSION)
            f.write(exc_text + "\n")
    except Exception:
        pass


def main():
    src_path = sys.argv[1] if len(sys.argv) > 1 else None
    if src_path and not os.path.exists(src_path):
        src_path = None

    root = tk.Tk()
    root.withdraw()          # 先不显示，等免责声明确认后再露面

    try:
        # 1) 先弹免责声明，确认后才继续
        show_disclaimer(root)

        # 2) 选总表原件（这个文件框挂在主窗口上，会显示在主界面前面）
        if not src_path:
            src_path = choose_file()
            if not src_path:
                root.destroy()
                return

        # 3) 显示主界面，并确保它在最前面、不会被别的窗口压住
        root.deiconify()
        ScoreApp(root, src_path)
        root.lift()
        root.focus_force()
        try:
            root.attributes("-topmost", True)
            root.after(400, lambda: root.attributes("-topmost", False))
        except Exception:
            pass
        root.mainloop()
    except Exception:
        import traceback
        detail = traceback.format_exc()
        write_crash_log(detail)
        try:
            messagebox.showerror(
                "程序出错",
                "程序启动/运行中出错，已在程序所在文件夹写了一份“出错日志.txt”，"
                "请把它发给作者（QQ %s）反馈。\n\n%s" % (AUTHOR_QQ, detail[-800:]))
        except Exception:
            pass
        raise


if __name__ == "__main__":
    main()
